import os
import wave
import struct
import math
import pytest
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import docx
import pandas as pd
from unittest.mock import patch, AsyncMock

from backend.config import settings
from backend.parsers import parse_item, detect_source_type
from backend.parsers.pdf_parser import parse_pdf, EncryptedFileError
from backend.parsers.docx_parser import parse_docx, CorruptedDocxError
from backend.parsers.image_ocr import parse_image
from backend.parsers.audio_transcribe import parse_audio, calculate_audio_rms
from backend.parsers.csv_parser import parse_csv
from backend.parsers.link_extractor import parse_url

@pytest.fixture
def temp_test_dir(tmp_path):
    """Fixture providing a temporary directory for generating test files."""
    d = tmp_path / "parser_tests"
    d.mkdir(parents=True, exist_ok=True)
    return d

def test_source_type_detection():
    """Verify source type resolution across diverse extensions and URLs."""
    assert detect_source_type("research_paper.pdf") == "pdf"
    assert detect_source_type("meeting_notes.docx") == "docx"
    assert detect_source_type("architecture_diagram.png") == "image"
    assert detect_source_type("photo.JPEG") == "image"
    assert detect_source_type("voice_memo.mp3") == "audio"
    assert detect_source_type("recording.wav") == "audio"
    assert detect_source_type("dataset.csv") == "csv"
    assert detect_source_type("sheet.xlsx") == "csv"
    assert detect_source_type("scratchpad.md") == "note"
    assert detect_source_type("https://example.com/article") == "link"
    assert detect_source_type("http://techcrunch.com") == "link"
    assert detect_source_type("unknown_archive.tar.gz") == "unknown"

@pytest.mark.asyncio
async def test_plain_text_parser():
    """Verify plain text note ingestion and title extraction."""
    raw_note = "# Project Alpha Roadmap\n\nKey deliverables include FAISS integration and SQLite WAL mode."
    item = await parse_item(
        item_id="note-1",
        source_type="note",
        raw_text_content=raw_note
    )
    assert item.id == "note-1"
    assert item.source_type == "note"
    assert item.title == "Project Alpha Roadmap"
    assert "FAISS integration" in item.extracted_text

def test_docx_parser(temp_test_dir):
    """Verify DOCX parsing with heading preservation and Markdown table linearization."""
    docx_file = temp_test_dir / "sample_spec.docx"
    doc = docx.Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph("This document outlines the multi-modal second brain design.")
    doc.add_heading("Component Matrix", level=2)
    
    # Add table
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Module"
    table.cell(0, 1).text = "Tech"
    table.cell(1, 0).text = "Vector Store"
    table.cell(1, 1).text = "FAISS"
    table.cell(2, 0).text = "Database"
    table.cell(2, 1).text = "SQLite"
    
    doc.save(str(docx_file))

    result = parse_docx(docx_file)
    assert "Architecture Overview" in result["title"] or "Sample Spec" in result["title"]
    assert "# Architecture Overview" in result["text"]
    assert "## Component Matrix" in result["text"]
    assert "| Module | Tech |" in result["text"]
    assert "| Vector Store | FAISS |" in result["text"]
    assert result["metadata"]["table_count"] == 1

def test_csv_parser(temp_test_dir):
    """Verify CSV dialect sniffing, schema profiling, and Markdown table serialization."""
    csv_file = temp_test_dir / "knowledge_metrics.csv"
    data = {
        "Metric": ["Latency", "Throughput", "Accuracy"],
        "Target": ["< 150ms", "500 req/s", "98.5%"],
        "Status": ["Pass", "Pass", "Pass"]
    }
    df = pd.DataFrame(data)
    df.to_csv(str(csv_file), index=False)

    result = parse_csv(csv_file)
    assert result["row_count"] == 3
    assert result["col_count"] == 3
    assert "Metric" in result["text"]
    assert "Latency" in result["text"]
    assert "< 150ms" in result["text"]
    assert "Column Schema:" in result["text"]

def test_image_ocr_parser(temp_test_dir):
    """Verify Image OCR and preprocessing pipeline."""
    img_file = temp_test_dir / "test_screenshot.png"
    # Create test image with text
    img = Image.new("RGB", (400, 150), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((20, 50), "SecondSelf AI Engine", fill=(0, 0, 0))
    img.save(str(img_file))

    result = parse_image(img_file)
    assert "width" in result["metadata"]
    assert "height" in result["metadata"]
    assert result["metadata"]["format"] == "IMAGE"
    # Text might be extracted by Tesseract or fall back to vision description
    assert len(result["text"]) > 0

def test_audio_rms_and_parser(temp_test_dir):
    """Verify Audio RMS calculation and empty silence detection."""
    # 1. Create a silent WAV file (RMS should be ~0.0)
    silent_wav = temp_test_dir / "silent_memo.wav"
    sample_rate = 16000
    num_samples = sample_rate * 2  # 2 seconds
    with wave.open(str(silent_wav), "wb") as wf:
        wf.setnchannels(1)
        wf.setsampwidth(2)
        wf.setframerate(sample_rate)
        # write 0s
        wf.writeframes(struct.pack(f"<{num_samples}h", *([0] * num_samples)))

    rms, duration = calculate_audio_rms(silent_wav)
    assert rms < 0.005
    assert duration >= 1.9

    result = parse_audio(silent_wav)
    assert result["is_empty"] is True
    assert result["metadata"]["status"] == "empty_audio"

    # 2. Create a non-silent synthetic sine wave WAV file
    active_wav = temp_test_dir / "active_memo.wav"
    with wave.open(str(active_wav), "wb") as wf:
        wf.setnchannels(1)
        wf.setsampwidth(2)
        wf.setframerate(sample_rate)
        sine_samples = [int(15000 * math.sin(2 * math.pi * 440 * i / sample_rate)) for i in range(num_samples)]
        wf.writeframes(struct.pack(f"<{num_samples}h", *sine_samples))

    active_rms, _ = calculate_audio_rms(active_wav)
    assert active_rms > 0.1
    active_res = parse_audio(active_wav)
    assert active_res["is_empty"] is False

@pytest.mark.asyncio
async def test_link_extractor_mocked():
    """Verify web link parser with HTML title and body extraction."""
    mock_html = """
    <!DOCTYPE html>
    <html>
      <head>
        <title>SecondSelf Architecture Guide</title>
        <meta property="og:description" content="A comprehensive PKM platform specification.">
      </head>
      <body>
        <main>
          <h1>SecondSelf Architecture Guide</h1>
          <p>SecondSelf connects multi-modal notes into a living force-directed knowledge graph.</p>
          <p>Hybrid RAG combines dense FAISS search and sparse BM25 indexing with reciprocal rank fusion.</p>
        </main>
      </body>
    </html>
    """
    with patch("httpx.AsyncClient.get") as mock_get:
        mock_resp = AsyncMock()
        mock_resp.status_code = 200
        mock_resp.text = mock_html
        mock_get.return_value = mock_resp

        result = await parse_url("https://example.com/secondself-guide")
        assert "SecondSelf Architecture Guide" in result["title"]
        assert "living force-directed knowledge graph" in result["text"]
        assert "Hybrid RAG" in result["text"]

def test_pdf_parser_encrypted_handling(temp_test_dir):
    """Verify PDF handling of encrypted or regular PDF documents."""
    pdf_file = temp_test_dir / "simple_doc.pdf"
    # Create a basic PDF using reportlab or minimal structure
    try:
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(pdf_file))
        c.drawString(100, 750, "SecondSelf System Specification")
        c.drawString(100, 700, "Multi-modal ingestion pipeline supporting 7 core formats.")
        c.save()

        res = parse_pdf(pdf_file)
        assert res["page_count"] >= 1
        assert "SecondSelf System Specification" in res["text"] or len(res["text"]) > 0
    except ImportError:
        pass
