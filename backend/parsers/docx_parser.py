import logging
import zipfile
from pathlib import Path
from typing import Dict, Any
import docx

logger = logging.getLogger(__name__)

class CorruptedDocxError(Exception):
    """Raised when a DOCX file is corrupted or not a valid OpenXML archive."""
    pass

def parse_docx(file_path: str | Path) -> Dict[str, Any]:
    """
    Linearization of OpenXML Word documents into clean, structured Markdown text,
    preserving heading hierarchy (#, ##), bullet lists, and tables.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {path}")

    try:
        doc = docx.Document(str(path))
    except (zipfile.BadZipFile, Exception) as e:
        logger.error(f"Failed to open DOCX '{path.name}': {e}")
        raise CorruptedDocxError(f"Corrupted or invalid DOCX archive: {path.name} ({e})")

    title = path.stem.replace("_", " ").replace("-", " ").title()
    content_parts = []
    paragraph_count = len(doc.paragraphs)
    table_count = len(doc.tables)

    # 1. Process paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        style_name = (p.style.name if p.style else "").lower()
        
        if "heading 1" in style_name:
            content_parts.append(f"# {text}")
        elif "heading 2" in style_name:
            content_parts.append(f"## {text}")
        elif "heading 3" in style_name:
            content_parts.append(f"### {text}")
        elif "heading" in style_name:
            content_parts.append(f"#### {text}")
        elif "list" in style_name or "bullet" in style_name:
            content_parts.append(f"- {text}")
        else:
            content_parts.append(text)

    # 2. Process and linearize tables into Markdown format
    for t_idx, table in enumerate(doc.tables):
        rows = table.rows
        if not rows:
            continue

        table_md_rows = []
        # Header row
        header_cells = [cell.text.strip().replace("\n", " ") for cell in rows[0].cells]
        table_md_rows.append("| " + " | ".join(header_cells) + " |")
        table_md_rows.append("| " + " | ".join(["---"] * len(header_cells)) + " |")

        # Data rows
        for row in rows[1:]:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            table_md_rows.append("| " + " | ".join(cells) + " |")

        table_md = "\n".join(table_md_rows)
        content_parts.append(f"\n{table_md}\n")

    full_text = "\n\n".join(content_parts).strip()
    if not full_text:
        full_text = f"[DOCX Document: {path.name} (Empty document)]"

    metadata = {
        "paragraph_count": paragraph_count,
        "table_count": table_count,
        "format": "DOCX"
    }

    return {
        "text": full_text,
        "title": title,
        "paragraph_count": paragraph_count,
        "metadata": metadata
    }
